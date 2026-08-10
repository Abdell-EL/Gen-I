from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts.ingest import build_article, extract_docx_blocks
from scripts.chunk import smart_chunk_article


class DocxTableIngestionTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = Path(tempfile.mkdtemp(prefix="docx-table-ingest-"))
        self.docx_path = self.tempdir / "table_fixture.docx"
        self._write_fixture(self.docx_path)

    def tearDown(self):
        shutil.rmtree(self.tempdir, ignore_errors=True)

    def _write_fixture(self, path: Path) -> None:
        doc = Document()
        doc.add_heading("Aide clôture SAV fibre", level=1)
        doc.add_paragraph("1. Règles métier explicites")
        table = doc.add_table(rows=4, cols=4)
        headers = ["Typologie", "Retail", "Wholesale", "Code unique"]
        row1 = ["Refait Branchement PB", "FTO DEF PB DIVERS", "FTO DEF PB DIVERS", "22"]
        row2 = ["Refait branchement PM", "FTO DEF PM PM", "FTO DEF PM PM", "23"]

        for col, value in enumerate(headers):
            table.cell(0, col).text = value

        for col, value in enumerate(row1):
            table.cell(1, col).text = value

        for col, value in enumerate(row2):
            table.cell(2, col).text = value

        table.cell(3, 0).text = ""
        table.cell(3, 1).text = ""
        table.cell(3, 2).text = ""
        table.cell(3, 3).text = ""

        doc.add_paragraph("Texte en français avec accents : clôture, défaillance, rétablissement.")
        doc.save(path)

    def test_extract_docx_blocks_preserves_table_relationships(self):
        blocks = extract_docx_blocks(self.docx_path)
        content = "\n\n".join(block.text for block in blocks)

        self.assertIn("Aide clôture SAV fibre", content)
        self.assertIn("Typologie", content)
        self.assertIn("Refait Branchement PB", content)
        self.assertIn("FTO DEF PB DIVERS", content)
        self.assertIn("clôture", content)
        self.assertEqual(sum(1 for block in blocks if block.block_type == "table"), 1)
        self.assertNotIn("|||", content)

    def test_build_article_and_chunking_keep_semantic_row_context(self):
        article = build_article(self.docx_path)
        chunks = smart_chunk_article(article)
        joined = "\n\n".join(chunk["text"] for chunk in chunks)

        self.assertIn("Refait Branchement PB", joined)
        self.assertIn("Retail", joined)
        self.assertIn("Wholesale", joined)
        self.assertIn("Code unique", joined)
        self.assertIn("FTO DEF PB DIVERS", joined)
        self.assertIn("FTO DEF PM PM", joined)
        self.assertIn("clôture", joined)
        self.assertNotIn("\n\n\n", joined)

    def test_multiple_rows_remain_distinguishable(self):
        article = build_article(self.docx_path)
        chunks = smart_chunk_article(article)
        target_chunks = [chunk["text"] for chunk in chunks if "Refait" in chunk["text"]]

        self.assertGreaterEqual(len(target_chunks), 2)
        self.assertTrue(any("Refait Branchement PB" in chunk for chunk in target_chunks))
        self.assertTrue(any("Refait branchement PM" in chunk for chunk in target_chunks))


if __name__ == "__main__":
    unittest.main()
